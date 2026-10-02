"""Rendu d'un `DeckSpec` en fichier .pptx (python-pptx), 16:9."""

import math
from io import BytesIO

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.presentation import Presentation as PptxPresentation
from pptx.slide import Slide
from pptx.util import Inches, Pt

from app.llm.presentations.schema import ChartSpec, DeckSpec, SlideSpec, TableSpec
from app.llm.presentations.themes import THEMES, Theme

_LAYOUT_TITLE, _LAYOUT_CONTENT, _LAYOUT_SECTION = 0, 1, 2
_LAYOUT_TWO_CONTENT, _LAYOUT_TITLE_ONLY = 3, 5

_SLIDE_W, _SLIDE_H = 13.333, 7.5
_MARGIN = 0.7
_BODY_TOP, _BODY_H = 1.85, 5.0
_BODY_W = _SLIDE_W - 2 * _MARGIN

_CHART_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _style_runs(paragraph, size: int, color: str, theme: Theme, bold: bool = False):
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = theme.font
        run.font.color.rgb = _rgb(color)


def _set_title(shape, text, *, size, color, theme, left, top, width, height, align):
    shape.left, shape.top = Inches(left), Inches(top)
    shape.width, shape.height = Inches(width), Inches(height)
    frame = shape.text_frame
    frame.text = text
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.paragraphs[0].alignment = align
    _style_runs(frame.paragraphs[0], size, color, theme, bold=True)


def _background(slide: Slide, color: str) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color)


def _bar(slide: Slide, color: str, left: float, top: float, width: float) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.07)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()


def _remove(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def _fit_font_size(texts: list[str], width_in: float, height_in: float) -> int:
    """Plus grande taille de police (pt) pour laquelle les puces tiennent dans la zone."""
    for size in (28, 24, 20, 18, 16, 14):
        chars_per_line = max(1.0, width_in * 72 / (size * 0.5))
        lines = sum(max(1, math.ceil(len(t) / chars_per_line)) for t in texts)
        needed = lines * size * 1.2 + len(texts) * size * 0.5
        if needed <= height_in * 72:
            return size
    return 12


def _fill_bullets(shape, bullets, theme, *, left, top, width, height) -> None:
    shape.left, shape.top = Inches(left), Inches(top)
    shape.width, shape.height = Inches(width), Inches(height)
    frame = shape.text_frame
    frame.word_wrap = True
    size = _fit_font_size(bullets, width, height)
    for index, text in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.space_after = Pt(size * 0.4)
        _style_runs(paragraph, size, theme.text_color, theme)


def _content_title(slide: Slide, spec: SlideSpec, theme: Theme) -> None:
    _background(slide, theme.background)
    _set_title(
        slide.shapes.title,
        spec.title,
        size=32,
        color=theme.title_color,
        theme=theme,
        left=_MARGIN,
        top=0.45,
        width=_BODY_W,
        height=1.0,
        align=PP_ALIGN.LEFT,
    )
    _bar(slide, theme.accent, _MARGIN, 1.5, 1.2)


def _add_marque(slide: Slide, theme: Theme) -> None:
    """Bloc "République Française" (texte + liseré tricolore) ; le logo officiel n'est pas embarqué."""
    left, top = 0.9, 0.6
    for index, color in enumerate(("000091", "FFFFFF", "E1000F")):
        stripe = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left + index * 0.3),
            Inches(top),
            Inches(0.3),
            Inches(0.1),
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = _rgb(color)
        stripe.line.color.rgb = _rgb("DDDDDD" if color == "FFFFFF" else color)
    box = slide.shapes.add_textbox(
        Inches(left - 0.1), Inches(top + 0.15), Inches(4), Inches(1.4)
    )
    frame = box.text_frame
    frame.word_wrap = True
    lines = (
        ("RÉPUBLIQUE", True),
        ("FRANÇAISE", True),
        ("Liberté Égalité Fraternité", False),
    )
    for index, (text, bold) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        _style_runs(paragraph, 16 if bold else 11, "161616", theme, bold=bold)


def _add_cover(prs: PptxPresentation, deck: DeckSpec, theme: Theme) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[_LAYOUT_TITLE])
    _background(slide, theme.cover_background)
    if theme.marque:
        _add_marque(slide, theme)
    _set_title(
        slide.shapes.title,
        deck.title,
        size=44,
        color=theme.cover_title_color,
        theme=theme,
        left=0.9,
        top=2.1,
        width=_SLIDE_W - 1.8,
        height=1.7,
        align=PP_ALIGN.LEFT,
    )
    _bar(slide, theme.accent, 0.9, 3.95, 1.5)
    subtitle = slide.placeholders[1]
    lines = [t for t in (deck.subtitle, deck.author) if t]
    if not lines:
        _remove(subtitle)
        return
    subtitle.left, subtitle.top = Inches(0.9), Inches(4.2)
    subtitle.width, subtitle.height = Inches(_SLIDE_W - 1.8), Inches(1.5)
    frame = subtitle.text_frame
    frame.word_wrap = True
    for index, text in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.alignment = PP_ALIGN.LEFT
        _style_runs(paragraph, 24 if index == 0 else 18, theme.cover_text_color, theme)


def _add_section(prs: PptxPresentation, spec: SlideSpec, theme: Theme) -> Slide:
    slide = prs.slides.add_slide(prs.slide_layouts[_LAYOUT_SECTION])
    _background(slide, theme.cover_background)
    _set_title(
        slide.shapes.title,
        spec.title,
        size=40,
        color=theme.cover_title_color,
        theme=theme,
        left=0.9,
        top=2.4,
        width=_SLIDE_W - 1.8,
        height=1.5,
        align=PP_ALIGN.LEFT,
    )
    _bar(slide, theme.accent, 0.9, 4.0, 1.5)
    body = slide.placeholders[1]
    if not spec.subtitle:
        _remove(body)
        return slide
    body.left, body.top = Inches(0.9), Inches(4.25)
    body.width, body.height = Inches(_SLIDE_W - 1.8), Inches(1.2)
    body.text_frame.word_wrap = True
    body.text_frame.text = spec.subtitle
    body.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    _style_runs(body.text_frame.paragraphs[0], 22, theme.cover_text_color, theme)
    return slide


def _add_bullets(prs: PptxPresentation, spec: SlideSpec, theme: Theme) -> Slide:
    slide = prs.slides.add_slide(prs.slide_layouts[_LAYOUT_CONTENT])
    _content_title(slide, spec, theme)
    _fill_bullets(
        slide.placeholders[1],
        spec.bullets,
        theme,
        left=_MARGIN,
        top=_BODY_TOP,
        width=_BODY_W,
        height=_BODY_H,
    )
    return slide


def _add_two_column(prs: PptxPresentation, spec: SlideSpec, theme: Theme) -> Slide:
    slide = prs.slides.add_slide(prs.slide_layouts[_LAYOUT_TWO_CONTENT])
    _content_title(slide, spec, theme)
    gap = 0.5
    width = (_BODY_W - gap) / 2
    for placeholder, bullets, left in (
        (slide.placeholders[1], spec.left, _MARGIN),
        (slide.placeholders[2], spec.right, _MARGIN + width + gap),
    ):
        _fill_bullets(
            placeholder,
            bullets,
            theme,
            left=left,
            top=_BODY_TOP,
            width=width,
            height=_BODY_H,
        )
    return slide


def _alt_text(graphic_frame, text: str) -> None:
    graphic_frame._element.nvGraphicFramePr.cNvPr.set("descr", text)


def _add_table(prs: PptxPresentation, spec: SlideSpec, theme: Theme) -> Slide:
    slide = prs.slides.add_slide(prs.slide_layouts[_LAYOUT_TITLE_ONLY])
    _content_title(slide, spec, theme)
    table_spec: TableSpec = spec.table
    n_rows, n_cols = len(table_spec.rows) + 1, len(table_spec.headers)
    size = 16 if n_rows <= 8 else 12
    row_h = max(0.32, min(0.55, _BODY_H / n_rows))
    frame = slide.shapes.add_table(
        n_rows,
        n_cols,
        Inches(_MARGIN),
        Inches(_BODY_TOP),
        Inches(_BODY_W),
        Inches(row_h * n_rows),
    )
    _alt_text(frame, f"Tableau : {spec.title}")
    table = frame.table
    for index in range(n_rows):
        table.rows[index].height = Inches(row_h)
    grid = [table_spec.headers, *table_spec.rows]
    for r, row in enumerate(grid):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = _rgb(theme.primary)
                color = theme.on_primary
            else:
                cell.fill.fore_color.rgb = _rgb(
                    theme.stripe if r % 2 == 0 else "FFFFFF"
                )
                color = theme.text_color
            _style_runs(cell.text_frame.paragraphs[0], size, color, theme, bold=r == 0)
    return slide


def _add_chart(prs: PptxPresentation, spec: SlideSpec, theme: Theme) -> Slide:
    slide = prs.slides.add_slide(prs.slide_layouts[_LAYOUT_TITLE_ONLY])
    _content_title(slide, spec, theme)
    chart_spec: ChartSpec = spec.chart
    data = CategoryChartData()
    data.categories = chart_spec.categories
    for series in chart_spec.series:
        data.add_series(series.name, series.values)

    frame = slide.shapes.add_chart(
        _CHART_TYPES[chart_spec.type],
        Inches(_MARGIN),
        Inches(_BODY_TOP),
        Inches(_BODY_W),
        Inches(_BODY_H),
        data,
    )
    _alt_text(frame, f"Graphique : {spec.title}")
    chart = frame.chart
    chart.font.size = Pt(14)
    chart.font.name = theme.font
    chart.has_title = False
    chart.has_legend = len(chart_spec.series) > 1 or chart_spec.type == "pie"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    plot = chart.plots[0]
    if chart_spec.type == "pie":
        plot.has_data_labels = True
        plot.data_labels.show_percentage = True
        plot.data_labels.show_value = False
        plot.data_labels.number_format = "0%"
        plot.data_labels.number_format_is_linked = False
        for index in range(len(chart_spec.categories)):
            point = plot.series[0].points[index]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(
                theme.series[index % len(theme.series)]
            )
        return slide

    for index, series in enumerate(plot.series):
        color = _rgb(theme.series[index % len(theme.series)])
        if chart_spec.type == "line":
            series.format.line.color.rgb = color
            series.format.line.width = Pt(3)
            series.smooth = False
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
    return slide


_ADDERS = {
    "bullets": _add_bullets,
    "two_column": _add_two_column,
    "table": _add_table,
    "chart": _add_chart,
}


def get_theme(name: str) -> Theme:
    theme = THEMES.get(name)
    if theme is None:
        raise ValueError(
            f"thème inconnu '{name}' ; thèmes disponibles : {', '.join(THEMES)}"
        )
    return theme


def build_pptx(deck: DeckSpec) -> bytes:
    theme = get_theme(deck.theme)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(_SLIDE_W), Inches(_SLIDE_H)
    prs.core_properties.title = deck.title
    prs.core_properties.author = deck.author

    _add_cover(prs, deck, theme)
    for spec in deck.slides:
        if spec.layout == "section":
            slide = _add_section(prs, spec, theme)
        else:
            slide = _ADDERS[spec.layout](prs, spec, theme)
        if spec.notes:
            slide.notes_slide.notes_text_frame.text = spec.notes

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
