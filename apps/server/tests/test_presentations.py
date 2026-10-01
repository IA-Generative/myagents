"""Tests for the presentation builder, signed download route and agent tool."""

import io
import time
import uuid
from unittest.mock import patch

import pytest
from langchain_core.messages import AIMessage
from pptx import Presentation
from pptx.util import Inches
from pydantic import ValidationError

from app.core import signing
from app.core.config import get_settings
from app.llm.agent_runtime import arun_agent_chat
from app.llm.client import LlmClient
from app.llm.presentations.builder import build_pptx
from app.llm.presentations.schema import DeckSpec
from app.llm.tools import list_presentation_themes
from app.schemas.agent import ChatMessage, ConfigSnapshot
from app.services import presentations as service
from tests.fakes import FakeToolCallingModel

DECK = {
    "title": "Bilan 2026",
    "subtitle": "Direction du numérique",
    "author": "Equipe produit",
    "slides": [
        {"layout": "section", "title": "Contexte", "subtitle": "Pourquoi maintenant"},
        {
            "layout": "bullets",
            "title": "Enjeux",
            "bullets": ["Accessibilité", "Sécurité"],
            "notes": "Insister sur la sécurité.",
        },
        {
            "layout": "two_column",
            "title": "Avant / Après",
            "left": ["Manuel"],
            "right": ["Automatisé"],
        },
        {
            "layout": "table",
            "title": "Budget",
            "table": {
                "headers": ["Poste", "Coût"],
                "rows": [["Dev", "10"], ["Ops", "5"]],
            },
        },
        {
            "layout": "chart",
            "title": "Évolution",
            "chart": {
                "type": "bar",
                "categories": ["2024", "2025"],
                "series": [{"name": "Usagers", "values": [10, 20]}],
            },
        },
        {
            "layout": "chart",
            "title": "Répartition",
            "chart": {
                "type": "pie",
                "categories": ["A", "B"],
                "series": [{"name": "Part", "values": [60, 40]}],
            },
        },
        {
            "layout": "chart",
            "title": "Tendance",
            "chart": {
                "type": "line",
                "categories": ["T1", "T2"],
                "series": [
                    {"name": "x", "values": [1, 2]},
                    {"name": "y", "values": [2, 1]},
                ],
            },
        },
    ],
}


@pytest.fixture(autouse=True)
def _signing_settings(monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "presentation_link_secret", "test-secret")
    monkeypatch.setattr(settings, "public_base_url", "http://test")


def test_build_pptx_renders_all_layouts():
    data = build_pptx(DeckSpec.model_validate(DECK))

    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == len(DECK["slides"]) + 1
    assert prs.slides[0].shapes.title.text == "Bilan 2026"
    assert (
        prs.slides[2].notes_slide.notes_text_frame.text == "Insister sur la sécurité."
    )
    assert any(s.has_table for s in prs.slides[4].shapes)
    assert all(any(s.has_chart for s in prs.slides[i].shapes) for i in (5, 6, 7))
    assert prs.slide_width == Inches(13.333)


def test_build_pptx_rejects_unknown_theme():
    with pytest.raises(ValueError, match="thème inconnu"):
        build_pptx(DeckSpec.model_validate({**DECK, "theme": "inexistant"}))


@pytest.mark.parametrize(
    "slide",
    [
        {"layout": "bullets", "title": "x"},
        {"layout": "two_column", "title": "x", "left": ["a"]},
        {"layout": "table", "title": "x"},
        {
            "layout": "table",
            "title": "x",
            "table": {"headers": ["a", "b"], "rows": [["1"]]},
        },
        {
            "layout": "chart",
            "title": "x",
            "chart": {
                "type": "pie",
                "categories": ["a"],
                "series": [
                    {"name": "s1", "values": [1]},
                    {"name": "s2", "values": [1]},
                ],
            },
        },
        {
            "layout": "chart",
            "title": "x",
            "chart": {
                "type": "bar",
                "categories": ["a", "b"],
                "series": [{"name": "s", "values": [1]}],
            },
        },
    ],
)
def test_slide_validation_rejects_incomplete_content(slide):
    with pytest.raises(ValidationError):
        DeckSpec.model_validate({"title": "t", "slides": [slide]})


def test_deck_limits_slide_count():
    slide = {"layout": "bullets", "title": "x", "bullets": ["a"]}
    with pytest.raises(ValidationError):
        DeckSpec.model_validate({"title": "t", "slides": [slide] * 41})


def test_signature_roundtrip_and_tampering():
    file_id = uuid.uuid4()
    exp = int(time.time()) + 60
    sig = signing._signature(file_id, exp)

    assert signing.verify(file_id, exp, sig)
    assert not signing.verify(file_id, exp + 1, sig)
    assert not signing.verify(uuid.uuid4(), exp, sig)
    assert not signing.verify(file_id, exp, "0" * 64)


def test_signature_expired():
    file_id = uuid.uuid4()
    exp = int(time.time()) - 1
    assert not signing.verify(file_id, exp, signing._signature(file_id, exp))


def test_signing_unavailable_without_secret(monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "presentation_link_secret", "")
    monkeypatch.setattr(settings, "openwebui_api_key", "")
    with pytest.raises(signing.SigningUnavailableError):
        signing.build_download_url(uuid.uuid4(), 60)


def test_list_presentation_themes_names_defaults():
    assert "institutionnel" in list_presentation_themes.invoke({})


async def _stored_file(db_session, ttl_minutes: int = 60):
    from datetime import timedelta

    return await service.save_file(
        db_session,
        filename="bilan.pptx",
        content_type=service.PPTX_CONTENT_TYPE,
        data=b"PKfake",
        ttl=timedelta(minutes=ttl_minutes),
    )


async def test_download_returns_file_with_valid_link(client, db_session):
    stored = await _stored_file(db_session)
    url = signing.build_download_url(stored.id, 60).removeprefix("http://test")

    res = await client.get(url)

    assert res.status_code == 200
    assert res.content == b"PKfake"
    assert res.headers["content-type"] == service.PPTX_CONTENT_TYPE
    assert 'filename="bilan.pptx"' in res.headers["content-disposition"]


async def test_download_rejects_bad_signature(client, db_session):
    stored = await _stored_file(db_session)
    exp = int(time.time()) + 60

    res = await client.get(
        f"/api/presentations/{stored.id}/download?exp={exp}&sig={'0' * 64}"
    )

    assert res.status_code == 403


async def test_download_rejects_expired_link(client, db_session):
    stored = await _stored_file(db_session)
    url = signing.build_download_url(stored.id, -10).removeprefix("http://test")

    assert (await client.get(url)).status_code == 403


async def test_download_unknown_file_is_404(client):
    file_id = uuid.uuid4()
    url = signing.build_download_url(file_id, 60).removeprefix("http://test")

    assert (await client.get(url)).status_code == 404


async def test_save_file_purges_expired(db_session):
    old = await _stored_file(db_session, ttl_minutes=-5)
    await _stored_file(db_session)

    assert await service.get_file(db_session, old.id) is None


async def test_agent_calls_create_presentation_tool(session_factory):
    config = ConfigSnapshot(
        name="Test",
        system_prompt="Tu crées des présentations.",
        tool_ids=["create_presentation"],
    )
    fake_model = FakeToolCallingModel(
        responses=[
            AIMessage(
                content="",
                tool_calls=[{"name": "create_presentation", "args": DECK, "id": "c1"}],
            ),
            AIMessage(content="Voici votre présentation."),
        ]
    )

    with (
        patch.object(LlmClient, "chat_model", return_value=fake_model),
        patch("app.llm.tools.SessionLocal", session_factory),
    ):
        reply = await arun_agent_chat(
            LlmClient(),
            config,
            history=[ChatMessage(role="user", content="Fais un bilan")],
            model="gpt-oss-120b",
            temperature=0.7,
        )

    assert reply == "Voici votre présentation."
    tool_message = fake_model.seen_messages[-1]
    assert "/api/presentations/" in tool_message.content
    assert "bilan-2026.pptx" in tool_message.content


async def test_open_webui_flow_returns_download_link(
    client, db_session, session_factory
):
    from app.api.deps import require_openwebui_key
    from app.main import app
    from app.scripts.seed_agents import DEFAULT_AGENTS, _agent_id, seed_default_agents

    app.dependency_overrides[require_openwebui_key] = lambda: None
    await seed_default_agents(db_session)
    slug = next(a["slug"] for a in DEFAULT_AGENTS if "presentation" in a["slug"])
    fake_model = FakeToolCallingModel(
        responses=[
            AIMessage(
                content="",
                tool_calls=[{"name": "create_presentation", "args": DECK, "id": "c1"}],
            ),
            AIMessage(content="Voici le lien."),
        ]
    )

    with (
        patch.object(LlmClient, "chat_model", return_value=fake_model),
        patch("app.llm.tools.SessionLocal", session_factory),
    ):
        res = await client.post(
            "/v1/chat/completions",
            json={
                "model": str(_agent_id(slug)),
                "messages": [{"role": "user", "content": "Fais un bilan"}],
            },
        )

    assert res.status_code == 200
    assert res.json()["choices"][0]["message"]["content"] == "Voici le lien."
    link = fake_model.seen_messages[-1].content.split("](")[1].rstrip(")")
    download = await client.get(link.removeprefix("http://test"))
    assert download.status_code == 200
    assert Presentation(io.BytesIO(download.content)).slides
