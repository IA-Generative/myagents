"""Helpers python-docx / python-pptx aux couleurs et à la typographie du DSFR.

Usage : from dsfr_office import new_docx, new_pptx
Marianne doit être installée sur le poste qui ouvre le fichier, sinon Word/PowerPoint retombe sur Arial.
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGB
from pptx.util import Inches, Pt as PptxPt

FONT = "Marianne"
BLUE_FRANCE = RGBColor(0x00, 0x00, 0x91)
RED_MARIANNE = RGBColor(0xE1, 0x00, 0x0F)
GREY_TITLE = RGBColor(0x16, 0x16, 0x16)
GREY_TEXT = RGBColor(0x3A, 0x3A, 0x3A)


def _style_font(style, size, color, bold=False):
    style.font.name = FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def new_docx(entity=""):
    """Document Word A4 : styles Marianne, titres bleu France, en-tête « RÉPUBLIQUE FRANÇAISE »."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(595.3), Pt(841.9)
    _style_font(doc.styles["Normal"], 10.5, GREY_TEXT)
    for name, size, color in (("Title", 26, BLUE_FRANCE), ("Heading 1", 20, BLUE_FRANCE), ("Heading 2", 15, BLUE_FRANCE), ("Heading 3", 12.5, GREY_TITLE)):
        _style_font(doc.styles[name], size, color, bold=True)

    head = sec.header.paragraphs[0]
    run = head.add_run("RÉPUBLIQUE FRANÇAISE")
    run.font.name, run.font.size, run.font.bold, run.font.color.rgb = FONT, Pt(10), True, GREY_TITLE
    if entity:
        head.add_run(f"\n{entity}").font.size = Pt(9)
    head.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return doc


def add_table(doc, header, rows):
    """Tableau DSFR : en-tête gris clair en gras, lignes sobres."""
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, header):
        cell.text = ""
        cell.paragraphs[0].add_run(str(label)).bold = True
        _shade(cell, "F6F6F6")
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = str(value)
    return table


def new_pptx():
    """Présentation 16:9 ; utiliser title_slide() et content_slide() pour garder la charte."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _text(slide, text, left, top, width, height, size, color, bold=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold = FONT, PptxPt(size), bold
    run.font.color.rgb = PptxRGB(*color)
    return tf


def title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _text(slide, "RÉPUBLIQUE FRANÇAISE", 0.8, 0.6, 6, 0.5, 14, (0x16, 0x16, 0x16), bold=True)
    _text(slide, title, 0.8, 2.6, 11.5, 1.6, 40, (0x00, 0x00, 0x91), bold=True)
    if subtitle:
        _text(slide, subtitle, 0.8, 4.4, 11.5, 1, 20, (0x16, 0x16, 0x16))
    return slide


def content_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _text(slide, title, 0.8, 0.5, 11.5, 1, 30, (0x00, 0x00, 0x91), bold=True)
    tf = _text(slide, bullets[0] if bullets else "", 0.8, 1.8, 11.5, 5, 20, (0x3A, 0x3A, 0x3A))
    for line in bullets[1:]:
        p = tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.name, run.font.size = FONT, PptxPt(20)
        run.font.color.rgb = PptxRGB(0x3A, 0x3A, 0x3A)
    return slide
